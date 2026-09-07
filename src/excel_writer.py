"""集計結果を書式付きExcelブックとして出力する"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
TITLE_FONT = Font(bold=True, size=13)
HEADER_FONT = Font(bold=True, size=10)
BASE_FONT = Font(size=10)

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NUM_FMT = "#,##0;[Red]-#,##0"
PCT_FMT = "0.0%;[Red]-0.0%"

MAX_WIDTH = 60


def _write_sheet(wb, title, subtitle, df, pct_columns=(), total_row_label=None):
    ws = wb.create_sheet(title)

    ws.cell(row=1, column=1, value=subtitle).font = TITLE_FONT
    header_row = 3

    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=header_row, column=j, value=col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, (_, row) in enumerate(df.iterrows(), start=header_row + 1):
        is_total = (
            total_row_label is not None
            and str(row.get(total_row_label[0], "")) == total_row_label[1]
        )
        for j, col in enumerate(df.columns, start=1):
            v = row[col]
            if v is not None and hasattr(v, "item"):
                v = v.item()
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(size=10, bold=is_total)
            c.border = BORDER
            if is_total:
                c.fill = TOTAL_FILL
            if col in pct_columns:
                c.number_format = PCT_FMT
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                c.number_format = NUM_FMT

    # 列幅の自動調整
    for j, col in enumerate(df.columns, start=1):
        longest = max(
            [len(str(col))]
            + [len(f"{v:,}") if isinstance(v, (int, float)) else len(str(v))
               for v in df[col].fillna("")]
        )
        ws.column_dimensions[get_column_letter(j)].width = min(longest + 3, MAX_WIDTH)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    return ws


def write_workbook(path, company, month, sheets):
    """sheets: [(シート名, 見出し, DataFrame, 率列のtuple, 合計行の指定), ...]"""
    wb = Workbook()
    wb.remove(wb.active)

    for name, subtitle, df, pct_cols, total_row in sheets:
        if df is None or df.empty:
            ws = wb.create_sheet(name)
            ws.cell(row=1, column=1, value=subtitle).font = TITLE_FONT
            ws.cell(row=3, column=1, value="該当データはありません").font = BASE_FONT
            continue
        _write_sheet(wb, name, subtitle, df, pct_cols, total_row)

    wb.save(path)
    return path
